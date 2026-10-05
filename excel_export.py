"""Exportación del listado visible a Excel con enlaces a Maximo Desktop."""
from __future__ import annotations

from pathlib import Path


def write_workorders_xlsx(path: str | Path, headers: list[str], rows: list[list[str]],
                          ot_column: int, uri_scheme: str = "maximodesk") -> int:
    """Escribe las filas recibidas preservando el orden visible en la interfaz."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from deep_links import make_ot_uri

    path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Órdenes de trabajo"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="17324D")
    header_font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28

    for row in rows:
        values = list(row)
        sheet.append(values)
        # Exportar como texto evita que contenido de Maximo que empiece por
        # «=» se interprete como una fórmula de Excel.
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
        cell = sheet.cell(sheet.max_row, ot_column + 1)
        ot = str(values[ot_column] or "").strip()
        if ot.isdecimal():
            cell.hyperlink = make_ot_uri(ot, development=uri_scheme == "maximodesk-dev")
            cell.style = "Hyperlink"

    if rows and headers:
        last_column = sheet.cell(1, len(headers)).column_letter
        table = Table(displayName="OrdenesTrabajo", ref=f"A1:{last_column}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)

    for column_index, header in enumerate(headers, start=1):
        width = min(48, max(12, len(str(header)) + 2))
        for row_index in range(2, sheet.max_row + 1):
            value = sheet.cell(row_index, column_index).value
            width = min(48, max(width, len(str(value or "")) + 2))
        sheet.column_dimensions[sheet.cell(1, column_index).column_letter].width = width

    workbook.save(path)
    return len(rows)
